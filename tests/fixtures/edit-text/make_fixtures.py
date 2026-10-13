"""Fixtures for the edit-text feature tests (requires reportlab, python-docx and LibreOffice).

Run: python3 tests/fixtures/edit-text/make_fixtures.py   (outputs are committed next to this file)
"""
import os
import shutil
import subprocess
import tempfile

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_RIGHT
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph

OUT = os.path.dirname(os.path.abspath(__file__))
DEJAVU = '/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf'

TEXT = ('The committee reviewed the proposal in detail and agreed to move forward with the plan. '
        'Funding will be released in two stages, subject to the usual quarterly reporting requirements '
        'and a final review at the end of the year.')


def subset_font():
    """Text in an embedded, subset TrueType font (only the glyphs it uses)."""
    pdfmetrics.registerFont(TTFont('DejaVuSans', DEJAVU))
    c = canvas.Canvas(os.path.join(OUT, 'subset-font.pdf'), pagesize=letter)
    c.setTitle('Subset font fixture')
    c.setFont('DejaVuSans', 16)
    c.drawString(72, 700, 'Embedded subset font sample')
    c.setFont('Helvetica', 12)
    c.drawString(72, 660, 'Plain Helvetica line')
    c.showPage()
    c.save()


def rotated():
    """A page with /Rotate 90 whose text reads upright on screen."""
    c = canvas.Canvas(os.path.join(OUT, 'rotated.pdf'), pagesize=letter)
    c.setTitle('Rotated page fixture')
    # Content is drawn sideways in user space so that it appears upright after /Rotate 90.
    c.saveState()
    c.translate(612, 0)
    c.rotate(90)
    # x now runs up the unrotated page (0..792) and y towards its left edge (0..612): once the
    # page is shown rotated clockwise this is an upright 792 x 612 landscape canvas
    c.setFont('Helvetica-Bold', 20)
    c.drawString(72, 520, 'Rotated heading')
    c.setFont('Helvetica', 12)
    c.drawString(72, 480, 'Text on a page that is rotated by ninety degrees.')
    c.restoreState()
    c.setPageRotation(90)
    c.showPage()
    c.save()


def layout_cases():
    c = canvas.Canvas(os.path.join(OUT, 'layout.pdf'), pagesize=letter)
    c.setTitle('Layout fixture')
    # 1. heading with descenders sitting tightly on a paragraph with accented capitals (1.0 em leading)
    c.setFont('Helvetica-Bold', 12)
    c.drawString(72, 720, 'Quarterly gyp jumping report')
    c.setFont('Times-Roman', 12)
    c.drawString(72, 708, 'ÉCOLE ÀNGSTRÖM ÔVERVIEW of results')
    c.drawString(72, 696, 'second line of the tight paragraph')
    c.drawString(72, 684, 'third line of the tight paragraph')

    # 2. justified paragraph
    style = ParagraphStyle('j', fontName='Times-Roman', fontSize=11, leading=14, alignment=TA_JUSTIFY)
    p = Paragraph(TEXT, style)
    w, h = p.wrap(300, 200)
    p.drawOn(c, 72, 640 - h)

    # 3. right aligned and centred paragraphs
    style = ParagraphStyle('r', fontName='Helvetica', fontSize=10, leading=13, alignment=TA_RIGHT)
    p = Paragraph('Right aligned text that wraps<br/>onto a second and shorter line', style)
    w, h = p.wrap(220, 100)
    p.drawOn(c, 320, 640 - h)
    style = ParagraphStyle('c', fontName='Helvetica', fontSize=10, leading=13, alignment=TA_CENTER)
    p = Paragraph('A centred paragraph of text<br/>with a short line', style)
    w, h = p.wrap(220, 100)
    p.drawOn(c, 320, 560 - h)

    # 4. an address block (hard line breaks that would otherwise fit)
    c.setFont('Helvetica', 11)
    for i, line in enumerate(['Jane Smith', '12 High Street', 'Springfield']):
        c.drawString(72, 470 - i * 14, line)

    # 5. linked text
    c.setFont('Helvetica', 11)
    c.drawString(72, 400, 'Read the full terms online')
    c.linkURL('https://example.org/terms', (72, 396, 220, 412), relative=0)

    # 6. a watermark crossing a paragraph
    c.setFont('Helvetica', 11)
    c.drawString(72, 330, 'Watermarked paragraph line one with plenty of words')
    c.drawString(72, 316, 'Watermarked paragraph line two with plenty of words')
    c.saveState()
    c.setFillColor(colors.Color(0.8, 0.1, 0.1))
    c.setFont('Helvetica-Bold', 30)
    c.translate(150, 300)
    c.rotate(25)
    c.drawString(0, 0, 'DRAFT')
    c.restoreState()

    # 7. coloured, mixed-style line
    c.setFont('Helvetica', 12)
    c.setFillColor(colors.HexColor('#C0392B'))
    c.drawString(72, 250, 'Red text and ')
    w = c.stringWidth('Red text and ', 'Helvetica', 12)
    c.setFont('Helvetica-Bold', 12)
    c.setFillColor(colors.black)
    c.drawString(72 + w, 250, 'bold words')

    c.setFont('Helvetica', 9)
    c.drawCentredString(306, 40, 'Footer centred on the page')
    c.showPage()
    c.save()


def shared_xobject():
    """Two pages drawing the same form XObject (a header), as page-mark tools produce."""
    c = canvas.Canvas(os.path.join(OUT, 'shared-xobject.pdf'), pagesize=letter)
    c.setTitle('Shared form XObject fixture')
    c.beginForm('Header')
    c.setFont('Helvetica-Bold', 14)
    c.drawString(72, 740, 'Shared header in a form XObject')
    c.endForm()
    for i in range(2):
        c.doForm('Header')
        c.setFont('Helvetica', 12)
        c.drawString(72, 700, 'Body text of page %d' % (i + 1))
        c.showPage()
    c.save()


def office_report():
    """A word-processor export: embedded subset TrueType fonts (Carlito/Caladea) with mixed styles."""
    from docx import Document

    tmp = tempfile.mkdtemp()
    try:
        d = Document()
        d.add_heading('Annual Report 2025', level=1)
        p = d.add_paragraph('This report summarises the activities of the year. ')
        r = p.add_run('Revenue grew by 12 percent')
        r.bold = True
        p.add_run(' while costs remained flat, which let the team invest in new tools and training for every '
                  'department in the company.')
        d.add_paragraph('Second paragraph with a café, naïve résumé and “smart quotes” — dash.')
        src = os.path.join(tmp, 'office-report.docx')
        d.save(src)
        subprocess.run(['soffice', '--headless', '--convert-to', 'pdf', '--outdir', tmp, src], check=True,
                       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        shutil.copy(os.path.join(tmp, 'office-report.pdf'), os.path.join(OUT, 'office-report.pdf'))
    finally:
        shutil.rmtree(tmp, ignore_errors=True)


if __name__ == '__main__':
    shared_xobject()
    office_report()
    subset_font()
    rotated()
    layout_cases()
